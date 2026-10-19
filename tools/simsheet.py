import os as _os
ROOT=_os.path.dirname(_os.path.dirname(_os.path.abspath(__file__)))
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.formatting.rule import CellIsRule, FormulaRule
P=ROOT+'/docs/KPI_Puntualita_Accuratezza.xlsx'
wb=load_workbook(P)
if 'Simulazione' in wb.sheetnames: del wb['Simulazione']
ws=wb.create_sheet('Simulazione')
F=lambda **k: Font(name='Arial',**k)
hdr_fill=PatternFill('solid',fgColor='16202E'); hdr_font=F(bold=True,color='FFFFFF')
sec_fill=PatternFill('solid',fgColor='DFE7EE')
blue=F(color='0000FF'); inp=PatternFill('solid',fgColor='FFF9E3')
thin=Side(style='thin',color='C3D0DC'); B=Border(bottom=thin)
EUR='€ #,##0;-€ #,##0;-'
def lab(r,c,t,bold=False,color=None,size=None):
    x=ws.cell(row=r,column=c,value=t); kw={'bold':bold}
    if color: kw['color']=color
    if size: kw['size']=size
    x.font=F(**kw); return x
def inv(r,c,v,fmt=None):
    x=ws.cell(row=r,column=c,value=v); x.font=blue; x.fill=inp
    if fmt: x.number_format=fmt
    return x
ws['A1']='Simulazione: Gino Bertolotti a Rivalta Scrivia, 24 mesi'; ws['A1'].font=F(bold=True,size=14)
ws['A2']='Modello mensile semplificato. Cambia le celle in blu (parametri, clienti, mesi di investimento) e la simulazione si ricalcola. Mese di inizio 0 = cliente o tecnologia esclusi.'; ws['A2'].font=F(color='4A5768')

# ---- parametri (A4:B..) ----
for c in range(1,3): ws.cell(row=4,column=c).fill=sec_fill
lab(4,1,'Parametri',True)
params=[('Cassa iniziale (€)',180000,EUR,'cassa0'),('Capienza magazzino (posti pallet)',2500,'#,##0','cap'),
('Tariffa stoccaggio (€/posto/mese)',5.5,'€ 0.00','tst'),('Tariffa movimentazione (€/pallet mosso)',3.8,'€ 0.00','tmo'),('Tariffa distribuzione (€/pallet consegnato)',30,'€ 0.00','tdi'),
('RAL magazziniere (€)',18500,EUR,'rmag'),('RAL autista (€)',20000,EUR,'raut'),('RAL impiegato (€)',22000,EUR,'rimp'),('Coefficiente costo azienda',1.42,'0.00','coef'),
('Compenso CEO (€/mese)',4000,EUR,'ceo'),('Immobile: manutenzione e utenze (€/mese)',3500,EUR,'imm'),('Spese varie (€/mese)',1500,EUR,'var'),
('Produttività magazziniere (pallet mossi/mese)',900,'#,##0','prod'),('Pallet per viaggio (motrice)',16,'0','ppv'),('Km per viaggio',240,'0','km'),('Viaggi al mese per autista',20,'0','vpa'),
('Prezzo gasolio (€/l)',0.86,'€ 0.00','fuel'),('Consumo motrice (l/km)',0.25,'0.00','cons'),('Pedaggi (€/km)',0.08,'€ 0.00','ped'),('Usura (€/km)',0.05,'€ 0.00','usu'),
('Costi fissi per motrice (€/mese)',1100,EUR,'mfix'),('Leasing motrice aggiuntiva (€/mese)',2100,EUR,'leas'),('Motrici di proprietà',3,'0','mprop'),
('Pallet consegnati/mese per impiegato oltre i 2 minimi',1500,'#,##0','ipv'),('Magazzinieri minimi',2,'0','minm'),('Impiegati minimi',2,'0','mini'),
('Puntualità di base (%)',92,'0.0','bP'),('Puntualità per punto di Operazioni sopra 10',0.4,'0.0','kP'),('Accuratezza di base (%)',96,'0.0','bA'),('Accuratezza per punto di Precisione sopra 10',0.3,'0.0','kA'),
('Assunzioni in ritardo di un mese (1 = sì, 0 = no)',0,'0','lag')]
R={}
for i,(t,v,fmt,key) in enumerate(params,5):
    lab(i,1,t); inv(i,2,v,fmt); R[key]=f'$B${i}'
pend=4+len(params)
# ---- personaggio e tecnologie (D4..) ----
for c in range(4,6): ws.cell(row=4,column=c).fill=sec_fill
lab(4,4,'Gino (CEO e manager di Rivalta)',True)
g=[('Operazioni',17,'op'),('Precisione',9,'prec'),('Leadership',14,'lead'),('Allergico al gestionale (1 = sì)',1,'all')]
for i,(t,v,key) in enumerate(g,5):
    lab(i,4,t); inv(i,5,v,'0'); R[key]=f'$E${i}'
for c in range(4,9): ws.cell(row=10,column=c).fill=sec_fill
lab(10,4,'Tecnologie',True); lab(10,5,'Mese di avvio',True); lab(10,6,'Accuratezza (pp)',True); lab(10,7,'Puntualità (pp)',True); lab(10,8,'Costo una tantum (€)',True); lab(10,9,'Canone (€/mese)',True); lab(10,10,'Settimane di apprendimento',True)
tech=[('Lettori barcode',3,1.5,0,12000,150,2),('WMS di base',0,1.5,0.5,65000,800,6)]
TR=[]
for i,(t,mese,a,p,una,can,sett) in enumerate(tech,11):
    lab(i,4,t); inv(i,5,mese,'0'); inv(i,6,a,'+0.0'); inv(i,7,p,'+0.0'); inv(i,8,una,EUR); inv(i,9,can,EUR); inv(i,10,sett,'0'); TR.append(i)
lab(13,4,'Durante l\'apprendimento la tecnologia toglie metà del suo effetto; "Allergico al gestionale" allunga l\'apprendimento del 50%.',color='4A5768')
# ---- clienti (D15..) ----
for c in range(4,11): ws.cell(row=15,column=c).fill=sec_fill
lab(15,4,'Clienti',True)
hdrs=['Cliente','Mese di inizio','Stock medio (posti)','Pallet consegnati/mese','Soglia puntualità (%)','Soglia accuratezza (%)']
for j,h in enumerate(hdrs,4):
    x=ws.cell(row=16,column=j,value=h); x.font=hdr_font; x.fill=hdr_fill; x.alignment=Alignment(wrap_text=True)
ws.row_dimensions[16].height=30
CL=[('Pastificio Mezzemaniche',1,700,300,92,None),('Caseificio Stracchinelli',4,600,480,93,96),('Ferramenta Bullonati',9,500,250,90,97),('Supermercati Carrellone',0,1500,2400,98,95)]
CR=[]
for i,(n,ini,st,pal,sp,sa) in enumerate(CL,17):
    lab(i,4,n,True); inv(i,5,ini,'0'); inv(i,6,st,'#,##0'); inv(i,7,pal,'#,##0'); inv(i,8,sp,'0.0'); inv(i,9,sa,'0.0'); CR.append(i)
lab(21,4,'Soglia accuratezza vuota = KPI non richiesto. Supermercati Carrellone è esclusa (mese 0): metti 4 per provare un contratto troppo grande.',color='4A5768')
lab(22,4,'Soddisfazione: parte da 60; +3 al mese sopra soglia, −4 se sotto di 2 punti al massimo, −10 e penale del 5% se oltre. Sotto 25 il cliente se ne va.',color='4A5768')

# ---- riepilogo (D24..) ----
for c in range(4,6): ws.cell(row=24,column=c).fill=sec_fill
lab(24,4,'Risultato',True)
# ---- griglia mensile ----
H0=pend+3; r0=H0+1; N=24; rN=r0+N-1
cols=['Mese','Stock (posti)','Saturazione','Pallet consegnati','Pallet mossi','Viaggi','Magazzinieri necessari','Autisti necessari','Impiegati necessari','Magazzinieri','Autisti','Impiegati','Produttività effettiva','Carico magazzinieri','Carico autisti','Effetto tecn. accuratezza','Effetto tecn. puntualità','Costo tecnologie (€)','Puntualità (%)','Accuratezza (%)','Morale','Ricavi (€)','Personale (€)','Mezzi (€)','Fissi (€)','Penali (€)','Utile (€)','Cassa (€)']
C={n:i+1 for i,n in enumerate(cols)}
start_cli=len(cols)+2
for k in range(4):
    base=start_cli+k*4
    for j,s in enumerate(['attivo','scarto KPI','soddisfazione','perso']):
        C[f'c{k}_{j}']=base+j
for name,ci in C.items():
    if name.startswith('c') and '_' in name:
        k=int(name[1]); j=int(name[3]); t=['Attivo','Scarto peggiore (pp)','Soddisfazione','Perso'][j]
        h=f'{CL[k][0]} – {t}'
    else: h=name
    x=ws.cell(row=H0,column=ci,value=h); x.font=hdr_font; x.fill=hdr_fill; x.alignment=Alignment(wrap_text=True,vertical='center')
ws.row_dimensions[H0].height=58
def c(n,r): return f'{L(C[n])}{r}'
for i in range(N):
    r=r0+i; p=r-1; first=(i==0)
    m=i+1
    ws.cell(row=r,column=C['Mese'],value=m)
    # client active flags
    act=[]; 
    for k in range(4):
        cr=CR[k]
        lost_prev='0' if first else c(f'c{k}_3',p)
        ws.cell(row=r,column=C[f'c{k}_0'],value=f'=IF(AND($E${cr}>0,{c("Mese",r)}>=$E${cr},{lost_prev}=0),1,0)')
    stock='+'.join([f'{c(f"c{k}_0",r)}*$F${CR[k]}' for k in range(4)])
    pal='+'.join([f'{c(f"c{k}_0",r)}*$G${CR[k]}' for k in range(4)])
    ws.cell(row=r,column=C['Stock (posti)'],value=f'={stock}')
    ws.cell(row=r,column=C['Saturazione'],value=f'={c("Stock (posti)",r)}/{R["cap"]}')
    ws.cell(row=r,column=C['Pallet consegnati'],value=f'={pal}')
    ws.cell(row=r,column=C['Pallet mossi'],value=f'=2*{c("Pallet consegnati",r)}')
    ws.cell(row=r,column=C['Viaggi'],value=f'=ROUNDUP({c("Pallet consegnati",r)}/{R["ppv"]},0)')
    ws.cell(row=r,column=C['Magazzinieri necessari'],value=f'=MAX({R["minm"]},ROUNDUP({c("Pallet mossi",r)}/{R["prod"]},0))')
    ws.cell(row=r,column=C['Autisti necessari'],value=f'=MAX(1,ROUNDUP({c("Viaggi",r)}/{R["vpa"]},0))')
    ws.cell(row=r,column=C['Impiegati necessari'],value=f'={R["mini"]}+INT({c("Pallet consegnati",r)}/{R["ipv"]})')
    for a,b in (('Magazzinieri','Magazzinieri necessari'),('Autisti','Autisti necessari'),('Impiegati','Impiegati necessari')):
        v=f'={c(b,r)}' if first else f'=IF({R["lag"]}=1,{c(b,p)},{c(b,r)})'
        ws.cell(row=r,column=C[a],value=v)
    mor_prev='65' if first else c('Morale',p)
    ws.cell(row=r,column=C['Produttività effettiva'],value=f'={R["prod"]}*(0.7+0.3*{mor_prev}/70)')
    ws.cell(row=r,column=C['Carico magazzinieri'],value=f'={c("Pallet mossi",r)}/({c("Magazzinieri",r)}*{c("Produttività effettiva",r)})')
    ws.cell(row=r,column=C['Carico autisti'],value=f'={c("Viaggi",r)}/({c("Autisti",r)}*{R["vpa"]})')
    # tech factor: IF(start>0 AND m>=start, IF((m-start+1)*4.3<=sett*IF(all,1.5,1),-0.5,1),0)
    fac=lambda tr: f'IF(AND($E${tr}>0,{c("Mese",r)}>=$E${tr}),IF(({c("Mese",r)}-$E${tr}+1)*4.3<=$J${tr}*IF({R["all"]}=1,1.5,1),-0.5,1),0)'
    ws.cell(row=r,column=C['Effetto tecn. accuratezza'],value='='+'+'.join([f'$F${tr}*{fac(tr)}' for tr in TR]))
    ws.cell(row=r,column=C['Effetto tecn. puntualità'],value='='+'+'.join([f'$G${tr}*{fac(tr)}' for tr in TR]))
    ws.cell(row=r,column=C['Costo tecnologie (€)'],value='='+'+'.join([f'IF(AND($E${tr}>0,{c("Mese",r)}=$E${tr}),$H${tr},0)+IF(AND($E${tr}>0,{c("Mese",r)}>=$E${tr}),$I${tr},0)' for tr in TR]))
    ws.cell(row=r,column=C['Puntualità (%)'],value=f'=MAX(50,MIN(99.9,{R["bP"]}+{R["kP"]}*({R["op"]}-10)+{c("Effetto tecn. puntualità",r)}-20*MAX(0,{c("Carico autisti",r)}-1)-5*MAX(0,{c("Carico magazzinieri",r)}-1.05)))')
    ws.cell(row=r,column=C['Accuratezza (%)'],value=f'=MAX(50,MIN(99.9,{R["bA"]}+{R["kA"]}*({R["prec"]}-10)+0.04*({mor_prev}-60)+{c("Effetto tecn. accuratezza",r)}-25*MAX(0,{c("Saturazione",r)}-0.85)-8*MAX(0,{c("Carico magazzinieri",r)}-1)))')
    ws.cell(row=r,column=C['Morale'],value=f'=MAX(0,MIN(100,{mor_prev}+0.35*(58+0.6*({R["lead"]}-10)-30*MAX(0,{c("Carico magazzinieri",r)}-1)-20*MAX(0,{c("Carico autisti",r)}-1)-{mor_prev})))')
    ws.cell(row=r,column=C['Ricavi (€)'],value=f'={c("Stock (posti)",r)}*{R["tst"]}+{c("Pallet mossi",r)}*{R["tmo"]}+{c("Pallet consegnati",r)}*{R["tdi"]}')
    ws.cell(row=r,column=C['Personale (€)'],value=f'=({c("Magazzinieri",r)}*{R["rmag"]}+{c("Autisti",r)}*{R["raut"]}+{c("Impiegati",r)}*{R["rimp"]})*{R["coef"]}/12')
    ws.cell(row=r,column=C['Mezzi (€)'],value=f'={c("Viaggi",r)}*{R["km"]}*({R["cons"]}*{R["fuel"]}+{R["ped"]}+{R["usu"]})+{c("Autisti",r)}*{R["mfix"]}+MAX(0,{c("Autisti",r)}-{R["mprop"]})*{R["leas"]}')
    ws.cell(row=r,column=C['Fissi (€)'],value=f'={R["ceo"]}+{R["imm"]}+{R["var"]}')
    # clients KPI
    pens=[]
    for k in range(4):
        cr=CR[k]
        ws.cell(row=r,column=C[f'c{k}_1'],value=f'=IF({c(f"c{k}_0",r)}=1,MIN({c("Puntualità (%)",r)}-$H${cr},IF($I${cr}="",9,{c("Accuratezza (%)",r)}-$I${cr})),"")')
        sprev='60' if first else c(f'c{k}_2',p)
        w=c(f'c{k}_1',r)
        ws.cell(row=r,column=C[f'c{k}_2'],value=f'=IF({c(f"c{k}_0",r)}=1,MAX(0,MIN(100,{sprev}+IF({w}>=0,3,IF({w}>=-2,-4,-10)))),{sprev})')
        lprev='0' if first else c(f'c{k}_3',p)
        ws.cell(row=r,column=C[f'c{k}_3'],value=f'=IF(OR({lprev}=1,AND({c(f"c{k}_0",r)}=1,{c(f"c{k}_2",r)}<25)),1,0)')
        pens.append(f'IF(AND({c(f"c{k}_0",r)}=1,{w}<-2),0.05*($F${cr}*{R["tst"]}+2*$G${cr}*{R["tmo"]}+$G${cr}*{R["tdi"]}),0)')
    ws.cell(row=r,column=C['Penali (€)'],value='='+'+'.join(pens))
    ws.cell(row=r,column=C['Utile (€)'],value=f'={c("Ricavi (€)",r)}-{c("Personale (€)",r)}-{c("Mezzi (€)",r)}-{c("Fissi (€)",r)}-{c("Costo tecnologie (€)",r)}-{c("Penali (€)",r)}')
    cprev=R['cassa0'] if first else c('Cassa (€)',p)
    ws.cell(row=r,column=C['Cassa (€)'],value=f'={cprev}+{c("Utile (€)",r)}')
    for name,ci in C.items():
        x=ws.cell(row=r,column=ci); x.font=F(bold=name in ('Utile (€)','Cassa (€)')); x.border=B
        if name in ('Saturazione','Carico magazzinieri','Carico autisti'): x.number_format='0%'
        elif '(€)' in name: x.number_format=EUR
        elif name in ('Puntualità (%)','Accuratezza (%)','Produttività effettiva','Morale','Effetto tecn. accuratezza','Effetto tecn. puntualità') or name.endswith('_1') or name.endswith('_2'): x.number_format='0.0'
        elif name in ('Stock (posti)','Pallet consegnati','Pallet mossi'): x.number_format='#,##0'
# summary formulas
cas=f'{L(C["Cassa (€)"])}{r0}:{L(C["Cassa (€)"])}{rN}'; ut=f'{L(C["Utile (€)"])}{r0}:{L(C["Utile (€)"])}{rN}'; mes=f'{L(C["Mese"])}{r0}:{L(C["Mese"])}{rN}'
summ=[('Cassa minima (€)',f'=MIN({cas})',EUR),('Cassa a fine simulazione (€)',f'={L(C["Cassa (€)"])}{rN}',EUR),
('Mese in cui la cassa va sotto zero',f'=IFERROR(INDEX({mes},MATCH(TRUE,INDEX({cas}<0,0),0)),"mai")','0'),
('Utile medio degli ultimi 6 mesi (€/mese)',f'=AVERAGE({L(C["Utile (€)"])}{rN-5}:{L(C["Utile (€)"])}{rN})',EUR),
('Clienti persi',f'='+'+'.join([f'{L(C[f"c{k}_3"])}{rN}' for k in range(4)]),'0')]
for i,(t,f,fmt) in enumerate(summ,25):
    lab(i,4,t); x=ws.cell(row=i,column=5,value=f); x.font=F(bold=True); x.number_format=fmt
# conditional formatting
ws.conditional_formatting.add(f'{L(C["Cassa (€)"])}{r0}:{L(C["Cassa (€)"])}{rN}',CellIsRule(operator='lessThan',formula=['0'],fill=PatternFill('solid',fgColor='F8D3D4')))
ws.conditional_formatting.add(f'{L(C["Utile (€)"])}{r0}:{L(C["Utile (€)"])}{rN}',CellIsRule(operator='lessThan',formula=['0'],font=Font(name='Arial',color='C62828',bold=True)))
ws.conditional_formatting.add(f'{L(C["Utile (€)"])}{r0}:{L(C["Utile (€)"])}{rN}',CellIsRule(operator='greaterThanOrEqual',formula=['0'],font=Font(name='Arial',color='23803A',bold=True)))
for k in range(4):
    col=L(C[f'c{k}_3'])
    ws.conditional_formatting.add(f'{col}{r0}:{col}{rN}',CellIsRule(operator='equal',formula=['1'],fill=PatternFill('solid',fgColor='F8D3D4')))
ws.column_dimensions['A'].width=44; ws.column_dimensions['B'].width=12; ws.column_dimensions['C'].width=3
ws.column_dimensions['D'].width=40
for ci in range(5,60): 
    if ci not in (1,2,3,4): ws.column_dimensions[L(ci)].width=12
ws.column_dimensions['A'].width=44
ws.freeze_panes=ws.cell(row=r0,column=2)
order=['KPI','Clienti','Regole','Personaggi','Talenti','Recruiting','Tecnologie','Eventi','Simulazione']
wb._sheets=[wb[n] for n in order]
wb.save(P); print('ok',r0,rN,C['Cassa (€)'])
