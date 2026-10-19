import os as _os
ROOT=_os.path.dirname(_os.path.dirname(_os.path.abspath(__file__)))
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
P=ROOT+'/docs/KPI_Puntualita_Accuratezza.xlsx'
wb=load_workbook(P)
if 'Eventi' in wb.sheetnames: del wb['Eventi']
ws=wb.create_sheet('Eventi')
F=lambda **k: Font(name='Arial',**k)
hdr_fill=PatternFill('solid',fgColor='16202E'); hdr_font=F(bold=True,color='FFFFFF')
thin=Side(style='thin',color='C3D0DC'); B=Border(bottom=thin)
ws['A1']='Catalogo eventi'; ws['A1'].font=F(bold=True,size=14)
ws['A2']='Ogni riga è un evento. Effetti indicativi da bilanciare. "Scelta speciale" compare solo se un personaggio ha il talento o l\'abilità indicati. Tono: Ironico / Serio.'; ws['A2'].font=F(color='4A5768')
cols=['ID','Categoria','Fonte','Titolo','Testo','Quando può accadere','Dal','Al','Frequenza','Scelta A','Effetti A','Scelta B','Effetti B','Scelta speciale','Requisito','Effetti speciale','Tono','Evento successivo']
for i,c in enumerate(cols,1):
    x=ws.cell(row=4,column=i,value=c); x.font=hdr_font; x.fill=hdr_fill; x.alignment=Alignment(wrap_text=True,vertical='center')
ws.row_dimensions[4].height=30
S24='Il Sole 24 Pallet'; GM='La Gazzetta del Muletto'; RT='Radio Traffico Nord-Ovest'; MM='Il Mattinale del Ministero'
E=[
('MET-01','Meteo e strade',RT,'Nebbia in Val Padana','Visibilità sotto i 50 metri da Pavia a Brescia. Gli autisti chiedono di rallentare.','Ottobre–febbraio; magazzino in Lombardia o Piemonte',2002,2040,'Comune','Rallenta tutto','Puntualità −6 pp per 7 giorni in Lombardia e Piemonte','Si viaggia lo stesso','Puntualità −2 pp; 15% di rischio incidente (−10 credibilità, mezzo fermo 30 giorni)','"Conosco ogni curva"','Talento: Un milione di km al volante','Puntualità −1 pp, nessun rischio','Ironico',''),
('MET-02','Meteo e strade',RT,'Neve al Traforo del Frejus','Il traforo chiude per neve. Le consegne per Lione e Parigi sono ferme.','Dicembre–marzo; flussi verso la Francia',2002,2040,'Comune','Devia dal Monte Bianco','+450 € per viaggio, puntualità −3 pp sui flussi esteri','Aspetta la riapertura','Puntualità −12 pp sui flussi esteri per 3 giorni','Prenota il treno (autostrada ferroviaria)','Innovazione ≥ 12','+200 € per viaggio, puntualità invariata','Ironico',''),
('MET-03','Meteo e strade',RT,'Allerta rossa in Liguria','Piogge fortissime: A10 e A6 a rischio chiusura, la Protezione Civile invita a non circolare.','Settembre–novembre; magazzino in provincia di Genova',2002,2040,'Rara','Chiudi il magazzino di Genova per 2 giorni','Ricavi −2 giorni, morale +5, nessun rischio','Resta aperto','30% di rischio danni al magazzino (−40.000 €) e infortuni (−8 credibilità)','','','','Serio','MET-03b Raccolta fondi per i comuni colpiti'),
('MET-04','Meteo e strade',RT,'Afa record','40 °C a Milano. I gruppi frigo vanno in crisi e in magazzino si respira a fatica.','Luglio–agosto',2002,2040,'Comune','Compra ventilatori e acqua per tutti','−3.000 €, morale +5','Turni anticipati all\'alba','Produttività −5% per 10 giorni, morale invariato','Pausa gelato offerta dal CEO','Talento: Cena col buyer','−800 €, morale +8, stress CEO −5','Ironico',''),
('MET-05','Meteo e strade',GM,'Mucche sulla statale','In Valle d\'Aosta una mandria blocca un camion per tre ore. L\'autista posta la foto: 40.000 condivisioni.','Maggio–settembre; flussi verso Aosta',2002,2040,'Rara','Pubblica la foto sui canali aziendali','Reputazione +3','Lascia perdere','Nessun effetto','','','','Ironico',''),
('POL-01','Politica e burocrazia',MM,'ZTL "Zero Camion"','Il nuovo assessore vieta i furgoni sotto Euro 4 nel centro di Milano dal mese prossimo.','Dal 2008; flussi verso Milano città',2008,2040,'Comune','Rinnova i furgoni','Costo dei nuovi mezzi; credibilità +3','Ricorso al TAR','−5.000 € di avvocati; 40% di vincere e rinviare il divieto di 1 anno','Una telefonata in Comune','Amministrazione ≥ 14','Proroga di 6 mesi, ma l\'assessore si ricorderà di te','Ironico','POL-01b L\'assessore chiede un favore'),
('POL-02','Politica e burocrazia',MM,'Decreto Semplificazione','Un decreto di 340 pagine semplifica le pratiche di trasporto. Nessuno ha capito cosa cambi.','Una volta ogni 2-3 anni',2002,2040,'Comune','Studiatelo con calma','Amministrazione rallentata per 30 giorni (costi +3%)','Assumi un consulente','−6.000 €, nessun rallentamento','"L\'ho già letto stanotte"','Talento: Excel-dipendente','Nessun rallentamento, stress CEO +5','Ironico',''),
('POL-03','Politica e burocrazia',S24,'Sciopero nazionale degli autotrasportatori','Le associazioni di categoria fermano i camion per 5 giorni contro il caro gasolio.','Prezzo del gasolio in forte aumento',2002,2040,'Rara','Aderisci','Puntualità −20 pp per 5 giorni, morale +10, credibilità +3','Lavora lo stesso','Morale −8; 25% di blocchi ai caselli (mezzi fermi 2 giorni)','Media tra sindacati e clienti','Leadership ≥ 14','Puntualità −8 pp, morale +5','Ironico',''),
('POL-04','Politica e burocrazia',MM,'Bando Fabbrica 4.0','Il bando rimborsa il 40% degli investimenti in tecnologia. La domanda è di 60 pagine.','Dal 2017, una volta l\'anno',2017,2040,'Comune','Partecipa','3 impiegati occupati per 30 giorni; rimborso dopo 6-18 mesi','Lascia perdere','Nessun effetto','Delegalo al Direttore Finanza','Amministrazione ≥ 12 del Direttore Finanza','Nessun impiegato occupato','Ironico','POL-04b Il rimborso arriva (forse)'),
('POL-05','Politica e burocrazia',GM,'Elezioni comunali','Entrambi i candidati sindaci ti chiedono un contributo per la campagna.','Ogni 5 anni nelle province con un tuo magazzino',2002,2040,'Comune','Sostieni il candidato A','−5.000 €; se vince, permessi più rapidi; se perde, controlli più frequenti','Resta neutrale','Nessun effetto','Sostieni entrambi','Negoziazione ≥ 15','−10.000 €, amico di chiunque vinca, 20% di scandalo (−8 reputazione)','Ironico',''),
('POL-06','Politica e burocrazia',MM,'Controllo dell\'Ispettorato del Lavoro','Gli ispettori si presentano alle 7:00 in magazzino.','Più probabile se usi cooperative',2002,2040,'Comune','Collabora','Multa in base alle irregolarità (0-30.000 €)','Chiedi un rinvio','Multa +20%, credibilità −3','"Abbiamo tutto in ordine"','Precisione ≥ 14 del manager del magazzino','Nessuna multa, credibilità +2','Ironico',''),
('LOC-01','Vita locale',S24,'Settimana della Moda','A Milano arrivano buyer da tutto il mondo. I clienti della moda raddoppiano gli ordini.','Febbraio e settembre; clienti del settore moda',2002,2040,'Comune','Straordinari per tutti','Ricavi moda ×2 per 10 giorni, morale −5','Usa i vettori esterni','Ricavi moda ×2, margine −30%','','','','Ironico',''),
('LOC-02','Vita locale',RT,'Passa il Giro d\'Italia','Una tappa blocca due province per un giorno intero.','Maggio',2002,2040,'Comune','Sponsorizza la tappa','−15.000 €, reputazione +6','Riorganizza le consegne','Puntualità −5 pp per 1 giorno','','','','Ironico',''),
('LOC-03','Vita locale',GM,'Derby della Mole','Metà dei magazzinieri di Torino chiede il pomeriggio libero per la partita.','Magazzino in provincia di Torino',2002,2040,'Comune','Concedi il pomeriggio','Produttività −40% per mezza giornata, morale +6','Nega il permesso','Morale −8; 30% di malattie improvvise','Megaschermo in magazzino','Leadership ≥ 12','−1.500 €, morale +10, produttività −10% per mezza giornata','Ironico',''),
('LOC-04','Vita locale',GM,'Il prete benedice i camion','Il parroco propone la benedizione annuale della flotta.','Magazzini fuori dalle grandi città',2002,2040,'Rara','Accetta e fai una donazione','−500 €, morale +4, un autista ritira le dimissioni','Rifiuta gentilmente','Morale −2 tra gli autisti più anziani','','','','Ironico',''),
('LOC-05','Vita locale',S24,'Olimpiadi invernali di Torino','Torino ospita le Olimpiadi: servono trasporti per impianti, atleti e merci.','Gennaio 2006',2006,2006,'Datato','Partecipa alla gara dei trasporti olimpici','Gara speciale: contratto di 2 mesi molto redditizio','Ignora','Traffico +20% in provincia di Torino per 2 mesi','','','','Ironico',''),
('LOC-06','Vita locale',S24,'Expo Milano','Sei mesi di esposizione universale: i cantieri e poi i padiglioni chiedono logistica.','2014–2015',2014,2015,'Datato','Apri un presidio vicino a Rho','Nuovo cliente temporaneo, reputazione +5','Ignora','Domanda in provincia di Milano +10% per 6 mesi','','','','Ironico',''),
('PER-01','Personale',GM,'Vincita al Superenalotto','Un magazziniere vince 2 milioni e si licenzia sul gruppo WhatsApp aziendale.','Qualsiasi momento',2002,2040,'Rara','Fagli gli auguri','Perdi un magazziniere esperto, morale +3','Chiedigli un investimento','30% di ottenere 100.000 € di prestito a tasso zero','','','','Ironico',''),
('PER-02','Personale',GM,'Il cane del manager','Il manager di Novara porta in ufficio il suo bassotto, Ragù.','Qualsiasi magazzino',2002,2040,'Rara','Permettilo','Morale +5; 20% che qualcuno sia allergico (−3 morale amministrazione)','Vietalo','Opinione del manager −15','','','','Ironico',''),
('PER-03','Personale',GM,'Cena di Natale finita male','Il Direttore Commerciale e il capo traffico hanno litigato sul karaoke.','Dicembre',2002,2040,'Comune','Organizza un chiarimento','Stress CEO +3, rapporto ristabilito','Lascia che si arrangino','Puntualità −2 pp per 2 mesi','Brindisi di riconciliazione','Talento: Cena col buyer','Opinione di entrambi verso il CEO +10','Ironico',''),
('PER-04','Personale',GM,'L\'autista influencer','Un tuo autista ha 300.000 follower su TikTok con i video dalla cabina.','Dal 2020',2020,2040,'Rara','Fallo testimonial','Reputazione +6; 20% di video imbarazzante (−5)','Chiedigli di smettere','Opinione dell\'autista −20, 40% che si licenzi','','','','Ironico',''),
('PER-05','Personale',GM,'Furto di gasolio','Di notte qualcuno ha svuotato tre serbatoi nel piazzale.','Magazzini senza sorveglianza',2002,2040,'Comune','Installa telecamere','−8.000 €, furti futuri −80%','Fidati del guardiano (74 anni)','Altri furti possibili, morale +2','','','','Ironico',''),
('PER-06','Personale',GM,'Il gestionale si aggiorna da solo','Venerdì sera il gestionale ha installato un aggiornamento. Lunedì nessuno stampa le bolle.','Gestionale installato',2002,2040,'Comune','Chiama l\'assistenza','Accuratezza −4 pp per 2 giorni, −1.200 €','Torna alle bolle a mano','Accuratezza −8 pp per 2 giorni','"Lo sistemo io"','Innovazione ≥ 14','Nessun effetto, stress +3','Ironico',''),
('CLI-01','Clienti e mercato',S24,'Cambio di buyer','Esselonga ha un nuovo buyer: vuole rinegoziare tutto e adora i report in PowerPoint.','Clienti della GDO',2002,2040,'Comune','Accetta la rinegoziazione','Tariffa −4%, soddisfazione +10','Difendi le condizioni','30% di disdetta anticipata','Presentazione da 80 slide','Talento: Parla per acronimi','Tariffa invariata, soddisfazione +5','Ironico',''),
('CLI-02','Clienti e mercato',S24,'Black Friday','Gli ordini dei clienti dell\'elettronica triplicano per una settimana.','Fine novembre, dal 2015',2015,2040,'Comune','Straordinari','Ricavi ×3, morale −8, accuratezza −3 pp','Assumi interinali','−6.000 €, accuratezza −5 pp','','','','Ironico',''),
('CLI-03','Clienti e mercato',GM,'Matrimonio del figlio del cliente','Il titolare del Pastificio Rummolo ti invita al matrimonio del figlio, a Benevento.','Clienti con soddisfazione alta',2002,2040,'Rara','Vai con un regalo all\'altezza','−2.500 €, soddisfazione +15','Manda un biglietto','Soddisfazione −5','','','','Ironico',''),
('CLI-04','Clienti e mercato',S24,'Un cliente fallisce','Un tuo cliente porta i libri in tribunale. Le fatture degli ultimi 3 mesi non saranno pagate.','Clienti con finanze deboli',2002,2040,'Rara','Insinuati nel fallimento','Recuperi il 20% in 2 anni','Vendi la merce rimasta in magazzino','Recuperi il 40% subito, 20% di contestazioni legali','','','','Serio',''),
('CLI-05','Clienti e mercato',GM,'Ordine alle 17:58 di venerdì','"Ci serve per lunedì alle 6, è vitale."','Qualsiasi cliente',2002,2040,'Comune','Accetta','Straordinari −2.000 €, morale −4, soddisfazione +8','Rifiuta con cortesia','Soddisfazione −6','','','','Ironico',''),
('CLI-06','Clienti e mercato',S24,'Crisi finanziaria','Le banche chiudono i rubinetti, i volumi calano in tutti i settori.','Settembre 2008',2008,2008,'Datato','Taglia i costi','Personale −10%, credibilità −10, cassa protetta','Resisti senza licenziare','Cassa sotto pressione per 18 mesi, credibilità +5','','','','Serio',''),
('CLI-07','Clienti e mercato',S24,'Pandemia e lockdown','Il Paese si ferma. Alimentare, farmaceutico ed e-commerce esplodono, il resto crolla.','Marzo 2020',2020,2020,'Datato','Proteggi il personale (dispositivi, turni)','−20.000 €, morale +10, produttività −10%','Massima produzione','Produttività +10%, rischio contagi e assenze, credibilità −8','','','','Serio',''),
('CON-01','Concorrenti',S24,'Pranzo con Klaus Effizienz','Il CEO di Logistix Global ti invita a pranzo: "Solo per conoscerci".','Azienda in crescita',2002,2040,'Rara','Ascolta l\'offerta','Ricevi un\'offerta di acquisto','Declina','Opinione di Logistix −10','Bluffa su un accordo con un altro gruppo','Negoziazione ≥ 16','Logistix non attacca le tue province per 2 anni','Ironico','CON-01b L\'offerta formale'),
('CON-02','Concorrenti',GM,'Lite dai Fratelli Brambilla','I fratelli Brambilla litigano sull\'eredità. Il minore ti offre i suoi clienti se lo assumi.','Dopo la morte del fondatore dei Brambilla',2002,2040,'Rara','Assumilo come manager','Nuovo manager (Negoziazione alta) e 2 clienti; opinione dei Brambilla −30','Rifiuta','Opinione dei Brambilla +10','','','','Ironico',''),
('CON-03','Concorrenti',S24,'Consegne in 2 ore','Spedisci.io lancia le consegne in 2 ore a Milano. Due tuoi clienti chiedono perché tu non lo fai.','Dal 2014; clienti in provincia di Milano',2014,2040,'Comune','Avvia un progetto pilota','−30.000 €, sblocca il servizio Express','Spiega che non serve','Soddisfazione −5 per 2 clienti','','','','Ironico',''),
]
for r,ev in enumerate(E,5):
    for c,v in enumerate(ev,1):
        x=ws.cell(row=r,column=c,value=v if v!='' else None); x.font=F(bold=(c==4)); x.alignment=Alignment(wrap_text=True,vertical='top'); x.border=B
        if c in (7,8): x.number_format='0'
widths=(9,16,20,24,42,30,7,7,10,24,32,24,32,24,24,30,9,26)
for i,w in enumerate(widths):
    ws.column_dimensions[chr(65+i)].width=w
ws.freeze_panes='E5'
end=4+len(E)
ws.auto_filter.ref=f'A4:R{end}'
ws.conditional_formatting.add(f'Q5:Q{end}',FormulaRule(formula=['$Q5="Serio"'],fill=PatternFill('solid',fgColor='F8D3D4')))
ws.conditional_formatting.add(f'I5:I{end}',FormulaRule(formula=['$I5="Datato"'],fill=PatternFill('solid',fgColor='DCEBFA')))
ws.conditional_formatting.add(f'I5:I{end}',FormulaRule(formula=['$I5="Rara"'],fill=PatternFill('solid',fgColor='FFF0C2')))
# riepilogo
s=end+2
ws.cell(row=s,column=1,value='Riepilogo').font=F(bold=True,size=12)
cats=['Meteo e strade','Politica e burocrazia','Vita locale','Personale','Clienti e mercato','Concorrenti']
ws.cell(row=s+1,column=1,value='Categoria').font=F(bold=True); ws.cell(row=s+1,column=3,value='Eventi').font=F(bold=True)
for k,cn in enumerate(cats,s+2):
    ws.cell(row=k,column=1,value=cn).font=F(); ws.merge_cells(start_row=k,end_row=k,start_column=1,end_column=2)
    ws.cell(row=k,column=3,value=f'=COUNTIF($B$5:$B${end},A{k})').font=F()
t=s+2+len(cats)
ws.cell(row=t,column=1,value='Totale').font=F(bold=True); ws.cell(row=t,column=3,value=f'=SUM(C{s+2}:C{t-1})').font=F(bold=True)
wb.save(P); print(len(E))
