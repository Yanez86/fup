import os as _os
ROOT=_os.path.dirname(_os.path.dirname(_os.path.abspath(__file__)))
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation

wb=Workbook()
F=lambda **k: Font(name='Arial',**k)
hdr_fill=PatternFill('solid',fgColor='16202E'); hdr_font=F(bold=True,color='FFFFFF')
title_font=F(bold=True,size=14)
thin=Side(style='thin',color='C3D0DC'); B=Border(bottom=thin)
blue=F(color='0000FF'); wrap=Alignment(wrap_text=True,vertical='top')
inp_fill=PatternFill('solid',fgColor='FFF9E3')

def header(ws,row,cols):
    for i,c in enumerate(cols,1):
        cell=ws.cell(row=row,column=i,value=c); cell.font=hdr_font; cell.fill=hdr_fill
        cell.alignment=Alignment(wrap_text=True,vertical='center')

# --- Foglio Regole (prima per i riferimenti) ---
wr=wb.active; wr.title='Regole'
wr['A1']='Regole dei KPI'; wr['A1'].font=title_font
wr['A3']='Parametro'; wr['B3']='Valore'; wr['C3']='Note'
header(wr,3,['Parametro','Valore','Note'])
wr['A4']='Margine "leggermente sotto" (punti %)'; wr['B4']=0.02; wr['B4'].number_format='0.0%'; wr['B4'].font=blue; wr['B4'].fill=inp_fill
wr['C4']='Ipotesi di design: sotto la soglia fino a questo margine = "Leggermente sotto"; oltre = "Molto sotto".'
wr['A5']='Saturazione oltre cui aumentano gli errori'; wr['B5']=0.85; wr['B5'].number_format='0%'; wr['B5'].font=blue; wr['B5'].fill=inp_fill
wr['C5']='Ipotesi di design usata per l\'accuratezza (vedi foglio KPI).'
header(wr,7,['Esito','Effetto sul cliente',''])
rows=[('Sopra soglia','La soddisfazione sale; possibili bonus o ampliamenti del contratto.'),
('Leggermente sotto','La soddisfazione cala.'),
('Molto sotto','Penali economiche; se dura più mesi, rischio di mancato rinnovo.'),
('Non richiesto','Il KPI non è nel contratto: non incide sulla soddisfazione.')]
for i,(a,b) in enumerate(rows,8):
    wr.cell(row=i,column=1,value=a).font=F(bold=True); c=wr.cell(row=i,column=2,value=b); c.font=F(); c.alignment=wrap
wr.merge_cells('B8:C8');wr.merge_cells('B9:C9');wr.merge_cells('B10:C10');wr.merge_cells('B11:C11')
wr.column_dimensions['A'].width=40; wr.column_dimensions['B'].width=22; wr.column_dimensions['C'].width=70
for r in (4,5): wr.cell(row=r,column=1).font=F(); wr.cell(row=r,column=3).font=F(color='4A5768'); wr.cell(row=r,column=3).alignment=wrap

# --- Foglio KPI ---
wk=wb.create_sheet('KPI',0)
wk['A1']='KPI di servizio: Puntualità e Accuratezza'; wk['A1'].font=title_font
wk['A2']='L\'OTIF è diviso in due KPI distinti. Ogni cliente può richiederne uno, entrambi o con soglie diverse.'; wk['A2'].font=F(color='4A5768')
header(wk,4,['KPI','Cosa misura','Fattore che lo influenza','Dettaglio'])
data=[
('Puntualità (On Time)','% di consegne arrivate entro la data o la fascia oraria concordata',[
 ('Capacità di trasporto','Mezzi e autisti disponibili rispetto agli ordini'),
 ('Tratta','Distanza e tipo di percorso: città, autostrada, valichi alpini'),
 ('Eventi','Neve, scioperi, blocchi stradali, guasti'),
 ('Persone','Abilità Operazioni del manager e del Direttore Operativo'),
 ('Tecnologia','Telematica, pianificazione delle rotte'),
 ('Vettori esterni','Affidabilità dei padroncini e dei vettori in subvezione')]),
('Accuratezza (In Full)','% di ordini consegnati completi e corretti: articoli e quantità giusti, senza danni',[
 ('Magazzinieri','Morale ed esperienza'),
 ('Saturazione','Oltre la soglia del foglio Regole gli errori aumentano'),
 ('Tecnologia','Lettori barcode, WMS, pick-to-light, robot'),
 ('Persone','Abilità del manager del magazzino'),
 ('Tipo di merce','Fragile, refrigerata o a lotti è più difficile da gestire')])]
r=5
for kpi,desc,facts in data:
    start=r
    for f,d in facts:
        wk.cell(row=r,column=3,value=f).font=F(bold=True)
        c=wk.cell(row=r,column=4,value=d); c.font=F(); c.alignment=wrap
        for col in range(1,5): wk.cell(row=r,column=col).border=B
        r+=1
    wk.cell(row=start,column=1,value=kpi).font=F(bold=True)
    c=wk.cell(row=start,column=2,value=desc); c.font=F(); c.alignment=wrap
    wk.merge_cells(start_row=start,end_row=r-1,start_column=1,end_column=1)
    wk.merge_cells(start_row=start,end_row=r-1,start_column=2,end_column=2)
    wk.cell(row=start,column=1).alignment=wrap
for col,w in zip('ABCD',(24,42,26,60)): wk.column_dimensions[col].width=w

# --- Foglio Clienti ---
wc=wb.create_sheet('Clienti',1)
wc['A1']='Requisiti dei clienti'; wc['A1'].font=title_font
wc['A2']='Celle in blu su sfondo giallo: da modificare. Soglia vuota = KPI non richiesto. I valori "attuali" sono esempi per provare gli esiti.'; wc['A2'].font=F(color='4A5768')
cols=['Cliente','Settore','Soglia puntualità','Soglia accuratezza','Perché','Puntualità attuale (esempio)','Accuratezza attuale (esempio)','Esito puntualità','Esito accuratezza','Stato complessivo']
header(wc,4,cols)
clients=[('Esselonga','GDO freschi',0.98,0.95,'Gli scaffali vuoti al mattino sono un disastro',0.975,0.96),
('ElettroUnieuri','Elettronica di consumo',0.90,0.99,'Un frigo sbagliato o rotto costa caro',0.93,0.955),
('Pastificio Rummolo','Alimentare secco',0.92,None,'Pasta secca: conta solo che arrivi',0.95,0.97),
('FarmaSpa','Farmaceutico',0.97,0.995,'Lotti e scadenze tracciati: cliente difficilissimo',0.97,0.99)]
for i,(n,s,p,a,why,pa,aa) in enumerate(clients,5):
    vals=[n,s,p,a,why,pa,aa]
    for j,v in enumerate(vals,1):
        c=wc.cell(row=i,column=j,value=v); c.font=F(); c.alignment=Alignment(vertical='top',wrap_text=(j==5))
        if j in (3,4,6,7):
            c.number_format='0.0%'; c.font=blue; c.fill=inp_fill
    m='Regole!$B$4'
    wc.cell(row=i,column=8,value=f'=IF(C{i}="","Non richiesto",IF(F{i}>=C{i},"Sopra soglia",IF(F{i}>=C{i}-{m},"Leggermente sotto","Molto sotto")))').font=F()
    wc.cell(row=i,column=9,value=f'=IF(D{i}="","Non richiesto",IF(G{i}>=D{i},"Sopra soglia",IF(G{i}>=D{i}-{m},"Leggermente sotto","Molto sotto")))').font=F()
    wc.cell(row=i,column=10,value=f'=IF(OR(H{i}="Molto sotto",I{i}="Molto sotto"),"Rischio penali",IF(OR(H{i}="Leggermente sotto",I{i}="Leggermente sotto"),"Soddisfazione in calo","In regola"))').font=F(bold=True)
    for col in range(1,11): wc.cell(row=i,column=col).border=B
for col,w in zip('ABCDEFGHIJ',(20,22,14,14,44,16,16,18,18,22)): wc.column_dimensions[col].width=w
wc.row_dimensions[4].height=32
rng='H5:J20'
def cf(text,color,fontc='16202E'):
    wc.conditional_formatting.add(rng,FormulaRule(formula=[f'ISNUMBER(SEARCH("{text}",H5))'],fill=PatternFill('solid',fgColor=color),font=Font(name='Arial',color=fontc)))
cf('Sopra','D9F2DF'); cf('In regola','D9F2DF'); cf('Leggermente','FFF0C2'); cf('calo','FFF0C2'); cf('Molto','F8D3D4'); cf('penali','F8D3D4'); cf('Non richiesto','EEF2F5','4A5768')
wc.freeze_panes='B5'
wb.save(ROOT+'/docs/KPI_Puntualita_Accuratezza.xlsx')
