import os as _os
ROOT=_os.path.dirname(_os.path.dirname(_os.path.abspath(__file__)))
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import ColorScaleRule
P=ROOT+'/docs/KPI_Puntualita_Accuratezza.xlsx'
wb=load_workbook(P)
for n in ('Talenti','Personaggi'):
    if n in wb.sheetnames: del wb[n]
F=lambda **k: Font(name='Arial',**k)
hdr_fill=PatternFill('solid',fgColor='16202E'); hdr_font=F(bold=True,color='FFFFFF')
thin=Side(style='thin',color='C3D0DC'); B=Border(bottom=thin)
blue=F(color='0000FF'); inp=PatternFill('solid',fgColor='FFF9E3')
def header(ws,row,cols):
    for i,c in enumerate(cols,1):
        x=ws.cell(row=row,column=i,value=c); x.font=hdr_font; x.fill=hdr_fill; x.alignment=Alignment(wrap_text=True,vertical='center',horizontal='center' if i>1 else 'left')
    ws.row_dimensions[row].height=34
AB=['Operazioni','Precisione','Negoziazione','Leadership','Amministrazione','Innovazione']

# ---------- TALENTI ----------
wt=wb.create_sheet('Talenti')
wt['A1']='Talenti'; wt['A1'].font=F(bold=True,size=14)
wt['A2']='Modificatori sulle sei abilità (scala 0-20) e sulla tenuta allo stress. Celle in blu: da bilanciare.'; wt['A2'].font=F(color='4A5768')
header(wt,4,['Talento']+AB+['Tenuta allo stress','Tipo','Altri effetti'])
T=[
('Un milione di km al volante',4,0,0,1,0,-2,0,'Origine','Gli autisti lo rispettano: morale autisti +5'),
('Allergico al gestionale',0,0,0,0,0,-4,0,'Negativo','Curva di apprendimento +50% nel suo magazzino'),
('Dialetto stretto',0,0,-2,2,0,0,0,'Misto','Opinione degli operai +10, dei clienti multinazionali −10'),
('Laureato in Bocconi',-1,0,1,0,3,0,0,'Origine','Accesso a eventi con banche e investitori'),
('Papà ha sempre ragione',0,0,0,-2,0,0,-5,'Negativo','Il fondatore interviene negli eventi'),
('Sempre al telefono',0,-1,2,0,0,0,-5,'Misto','Risponde ai clienti anche in ferie'),
('Parla per acronimi',0,0,2,-3,0,0,0,'Misto','Bonus nelle gare dei clienti grandi'),
('Stacanovista',2,0,0,0,0,0,-15,'Misto','Lo stress sale il doppio'),
('Cena col buyer',0,0,4,0,0,0,-10,'Positivo','Scelte speciali negli eventi con i clienti'),
('Excel-dipendente',0,3,0,-2,2,0,0,'Misto','Report KPI più precisi'),
('Parla con i robot',0,0,0,-2,0,3,0,'Misto','Curva di apprendimento −30%'),
('Calma olimpica',0,0,0,1,0,0,20,'Positivo','Ignora gli eventi di panico'),
('Ha fatto il militare',0,1,0,3,0,0,5,'Misto','Opinione dei dipendenti giovani −10'),
('Ex sindacalista',0,0,1,3,0,0,0,'Positivo','Vertenze sindacali −50%'),
('Conosce tutti al bar',0,0,1,2,0,0,0,'Positivo','Scopre in anticipo le mosse dei concorrenti nella sua provincia'),
('Braccino corto',0,0,0,-2,3,0,0,'Misto','Costi del suo magazzino −5%, morale −5'),
('Mano larga',0,0,0,2,-2,0,0,'Misto','Morale +5, costi del suo magazzino +5%'),
('Ansia da prestazione',0,2,0,0,0,0,-20,'Misto','Precisione alta, crolla sotto pressione'),
('Insonne',-2,-2,0,0,0,0,-10,'Negativo','Arriva con lo stress alto; sparisce con le ferie'),
('Tifo sfegatato',0,0,0,1,0,0,0,'Misto','Eventi speciali nei giorni del derby'),
]
T0=5
for r,t in enumerate(T,T0):
    wt.cell(row=r,column=1,value=t[0]).font=F(bold=True)
    for c in range(1,8):
        x=wt.cell(row=r,column=c+1,value=t[c]); x.font=blue; x.fill=inp; x.number_format='+0;-0;0'; x.alignment=Alignment(horizontal='center')
    wt.cell(row=r,column=9,value=t[8]).font=F()
    x=wt.cell(row=r,column=10,value=t[9]); x.font=F(); x.alignment=Alignment(wrap_text=True,vertical='top')
    for c in range(1,11): wt.cell(row=r,column=c).border=B
T_END=T0+len(T)-1
wt.column_dimensions['A'].width=30
for col in 'BCDEFGH': wt.column_dimensions[col].width=13
wt.column_dimensions['I'].width=11; wt.column_dimensions['J'].width=58
wt.freeze_panes='B5'

# ---------- PERSONAGGI ----------
wp=wb.create_sheet('Personaggi')
wp['A1']='Personaggi'; wp['A1'].font=F(bold=True,size=14)
wp['A2']='Abilità di base + modificatori dei talenti = valore finale (limitato tra 0 e 20). I talenti si scelgono dall\'elenco del foglio Talenti. Celle in blu: da modificare.'; wp['A2'].font=F(color='4A5768')
cols=['Nome','Età','Gruppo','Ruolo','Provincia','Talento 1','Talento 2','Talento 3']+[f'{a} base' for a in AB]+[f'{a}' for a in AB]+['Tenuta allo stress','Biografia']
header(wp,4,cols)
C=[
('Gino Bertolotti',52,'CEO giocabile','Titolare, ex padroncino','Alessandria','Un milione di km al volante','Allergico al gestionale','Dialetto stretto',(13,9,7,11,6,8),'Trent\'anni sulle strade d\'Europa, poi il capannone di Rivalta Scrivia comprato con la liquidazione e un mutuo. Conosce ogni casello dell\'A7 e diffida di tutto ciò che ha una presa elettrica.'),
('Edoardo Ferretti',31,'CEO giocabile','Figlio del fondatore','Milano','Laureato in Bocconi','Papà ha sempre ragione','Sempre al telefono',(7,8,10,8,9,10),'Ha ereditato il magazzino di Pioltello e un padre che passa "solo per un caffè" tutti i giorni alle 8. Vuole dimostrare di valere più del cognome.'),
('Valeria Conti',44,'CEO giocabile','Ex manager di multinazionale','Novara','Parla per acronimi','Stacanovista','Cena col buyer',(11,10,12,8,10,9),'Quindici anni di slide e KPI in una multinazionale. Licenziata in una ristrutturazione, ha deciso di fare da sé. Dorme quattro ore a notte e lo considera un vantaggio competitivo.'),
('Marco Pellegrino',36,'CEO giocabile','Ingegnere gestionale','Genova','Excel-dipendente','Parla con i robot','Calma olimpica',(9,10,5,8,9,14),'Ha progettato magazzini automatici per altri, ora ne vuole uno suo. Il piccolo deposito di Voltri è il prototipo: ogni pallet ha un codice, ogni codice ha un foglio Excel.'),
('Rosa Brambati',47,'Manager disponibile','Manager di magazzino','Lodi','Ex sindacalista','Mano larga','',(12,11,6,12,7,6),'Ha iniziato sul muletto, ha guidato la RSU per dieci anni. I magazzinieri la seguirebbero ovunque, i commercialisti un po\' meno.'),
('Salvatore "Totò" Esposito',58,'Manager disponibile','Capo traffico','Milano','Conosce tutti al bar','Sempre al telefono','',(14,8,10,9,6,5),'Trasferito da Napoli nel 1979, conosce ogni padroncino della provincia per nome e soprannome. Ha tre cellulari e li usa tutti.'),
('Chiara Ferraris',29,'Manager disponibile','Manager junior','Torino','Excel-dipendente','Ansia da prestazione','',(9,13,7,6,11,11),'Laurea al Politecnico, due stage, un master. Prepara report impeccabili, ma un ritardo di un camion le rovina la settimana.'),
('Davide Rinaldi',39,'Manager disponibile','Manager di magazzino','Bergamo','Stacanovista','Braccino corto','',(13,10,7,8,9,8),'Ha portato un magazzino di Orio dal 70% al 99% di accuratezza. Nessuno sa come, nessuno ha avuto un aumento.'),
('Luca Mazzoleni',33,'Manager disponibile','Responsabile tecnologia','Brescia','Parla con i robot','Insonne','',(8,10,5,6,8,15),'Programma di notte, installa di giorno. Chiama i robot del magazzino per nome.'),
('Franca Gallo',61,'Manager disponibile','Responsabile amministrazione','Alessandria','Calma olimpica','Braccino corto','',(7,13,8,9,14,4),'Ha visto passare la lira, l\'euro e tre gestionali. Le fatture le controlla ancora a mano, una per una, e non sbaglia mai.'),
('Klaus Effizienz',55,'CEO concorrente','CEO di Logistix Global','Milano','Parla per acronimi','Calma olimpica','',(15,14,14,9,15,13),'Arrivato da Amburgo per "razionalizzare il mercato italiano". Trova la pausa caffè una perdita di efficienza.'),
('Ambrogio Brambilla',79,'CEO concorrente','Fondatore di F.lli Brambilla','Milano','Dialetto stretto','Braccino corto','',(12,9,10,13,11,3),'Fondò l\'azienda nel 1958 con un camion e un mulo. Non ha ancora deciso a quale dei tre figli lasciarla, e i figli non si parlano.'),
('Renato Bartolazzi',63,'CEO concorrente','CEO del Gruppo Bartolazzi','Fuori mappa','Conosce tutti al bar','Ha fatto il militare','',(14,10,11,12,10,7),'Guida da Bologna una rete di corrieri in tutta Italia. Nel Nord-Ovest difende ogni cliente come se fosse l\'ultimo.'),
('Denis Fatica',45,'CEO concorrente','Presidente della Cooperativa La Fatica','Pavia','Braccino corto','Conosce tutti al bar','',(10,6,11,7,12,5),'Vince le gare con i prezzi più bassi della regione. Come ci riesca è oggetto di tre indagini e di molte chiacchiere.'),
('Ludovica Sprint',27,'CEO concorrente','Fondatrice di Spedisci.io (dal 2012)','Milano','Parla con i robot','Sempre al telefono','Insonne',(7,8,13,9,6,16),'Ha raccolto 20 milioni dagli investitori con una presentazione di 12 slide. Promette consegne in due ore e un\'app per tutto.'),
]
P0=5
for r,ch in enumerate(C,P0):
    nome,eta,gr,ruolo,prov,t1,t2,t3,base,bio=ch
    vals=[nome,eta,gr,ruolo,prov,t1,t2,t3 or None]
    for c,v in enumerate(vals,1):
        x=wp.cell(row=r,column=c,value=v); x.font=F(bold=(c==1)); x.alignment=Alignment(vertical='top',wrap_text=True)
        if c in (6,7,8): x.font=blue; x.fill=inp
    for k,b in enumerate(base):
        x=wp.cell(row=r,column=9+k,value=b); x.font=blue; x.fill=inp; x.alignment=Alignment(horizontal='center',vertical='top')
    for k in range(6):
        bc=wp.cell(row=r,column=9+k).column_letter
        tc=chr(ord('B')+k)  # column in Talenti
        mods='+'.join([f'IFERROR(INDEX(Talenti!${tc}${T0}:${tc}${T_END},MATCH({t}{r},Talenti!$A${T0}:$A${T_END},0)),0)' for t in ('F','G','H')])
        x=wp.cell(row=r,column=15+k,value=f'=MAX(0,MIN(20,{bc}{r}+{mods}))'); x.font=F(bold=True); x.alignment=Alignment(horizontal='center',vertical='top')
    mods='+'.join([f'IFERROR(INDEX(Talenti!$H${T0}:$H${T_END},MATCH({t}{r},Talenti!$A${T0}:$A${T_END},0)),0)' for t in ('F','G','H')])
    x=wp.cell(row=r,column=21,value=f'=50+{mods}'); x.alignment=Alignment(horizontal='center',vertical='top'); x.font=F()
    x=wp.cell(row=r,column=22,value=bio); x.font=F(); x.alignment=Alignment(wrap_text=True,vertical='top')
    for c in range(1,23): wp.cell(row=r,column=c).border=B
P_END=P0+len(C)-1
dv=DataValidation(type='list',formula1=f'=Talenti!$A${T0}:$A${T_END}',allow_blank=True)
wp.add_data_validation(dv); dv.add(f'F{P0}:H{P_END+20}')
wp.conditional_formatting.add(f'O{P0}:T{P_END}',ColorScaleRule(start_type='num',start_value=0,start_color='F8D3D4',mid_type='num',mid_value=10,mid_color='FFFFFF',end_type='num',end_value=20,end_color='BFE8C9'))
widths={'A':24,'B':6,'C':18,'D':26,'E':12,'F':24,'G':24,'H':22,'U':11,'V':70}
for k,v in widths.items(): wp.column_dimensions[k].width=v
for k in range(9,21): wp.column_dimensions[wp.cell(row=4,column=k).column_letter].width=11
wp.freeze_panes='B5'
order=['KPI','Clienti','Regole','Personaggi','Talenti','Recruiting','Tecnologie','Eventi']
wb._sheets=[wb[n] for n in order]
wb.save(P)
