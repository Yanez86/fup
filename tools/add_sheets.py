import os as _os
ROOT=_os.path.dirname(_os.path.dirname(_os.path.abspath(__file__)))
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
P=ROOT+'/docs/KPI_Puntualita_Accuratezza.xlsx'
wb=load_workbook(P)
F=lambda **k: Font(name='Arial',**k)
hdr_fill=PatternFill('solid',fgColor='16202E'); hdr_font=F(bold=True,color='FFFFFF')
sec_fill=PatternFill('solid',fgColor='DFE7EE')
title_font=F(bold=True,size=14)
thin=Side(style='thin',color='C3D0DC'); B=Border(bottom=thin)
blue=F(color='0000FF'); inp=PatternFill('solid',fgColor='FFF9E3')
wrap=Alignment(wrap_text=True,vertical='top')
EUR='€ #,##0;-€ #,##0;-'
def header(ws,row,cols,c0=1):
    for i,c in enumerate(cols,c0):
        x=ws.cell(row=row,column=i,value=c); x.font=hdr_font; x.fill=hdr_fill; x.alignment=Alignment(wrap_text=True,vertical='center')
    ws.row_dimensions[row].height=32
def section(ws,row,text,ncol):
    c=ws.cell(row=row,column=1,value=text); c.font=F(bold=True,size=12)
    for i in range(1,ncol+1): ws.cell(row=row,column=i).fill=sec_fill
def put(ws,r,c,v,fmt=None,inp_=False,bold=False,wrapit=False):
    x=ws.cell(row=r,column=c,value=v); x.font=blue if inp_ else F(bold=bold)
    if inp_: x.fill=inp
    if fmt: x.number_format=fmt
    x.alignment=Alignment(wrap_text=wrapit,vertical='top'); x.border=B
    return x

# ================= RECRUITING =================
ws=wb.create_sheet('Recruiting')
ws['A1']='Recruiting: compensi, welfare, canali e credibilità'; ws['A1'].font=title_font
ws['A2']='Valori di partenza indicativi per il 2002, da bilanciare. Celle in blu su sfondo giallo: da modificare. Il resto sono formule.'; ws['A2'].font=F(color='4A5768')

# 1. Ruoli
section(ws,4,'1. Ruoli e stipendio di mercato (RAL di riferimento, anno 2002)',8)
header(ws,5,['Ruolo','RAL di mercato (€)','Costo azienda (€/anno)','Disponibilità candidati (1-5)','Statistica principale','Note'])
roles=[('Autista patente CE',23000,2,'Operazioni','Carenza cronica: pochi candidati, molto contesi'),
('Autista patente C (furgoni/motrici)',20000,3,'Operazioni',''),
('Magazziniere / mulettista',18500,4,'Operazioni','Spesso fornito da cooperative'),
('Impiegato traffico / customer service',22000,3,'Commerciale','Pianifica i viaggi, parla con i clienti'),
('Capo turno magazzino',27000,3,'Persone',''),
('Manager di magazzino',48000,2,'Operazioni/Persone','Il "vassallo" che governa un magazzino'),
('Direttore (consiglio di sede)',75000,1,'Dipende dal ruolo','Solo tramite head hunter o promozione')]
put(ws,4,8,None)
ws['H5']='Coefficiente costo azienda'; ws['H5'].font=hdr_font; ws['H5'].fill=hdr_fill; ws['H5'].alignment=Alignment(wrap_text=True)
put(ws,6,8,1.42,'0.00',inp_=True)
ws['H7']='Ipotesi: contributi, TFR e oneri ≈ +42% sulla RAL'; ws['H7'].font=F(color='4A5768',size=9); ws['H7'].alignment=wrap
for i,(n,ral,d,stat,note) in enumerate(roles,6):
    put(ws,i,1,n,bold=True); put(ws,i,2,ral,EUR,True); put(ws,i,3,f'=B{i}*$H$6',EUR); put(ws,i,4,d,'0',True); put(ws,i,5,stat); put(ws,i,6,note,wrapit=True)
R_END=5+len(roles)

# 2. Province
r=R_END+2
section(ws,r,'2. Province del Nord-Ovest: coefficiente stipendi e bacino di candidati',8)
header(ws,r+1,['Provincia','Coeff. stipendi','Bacino candidati (1-5)','Concorrenza (1-5)','Polo logistico','RAL autista CE in provincia (€)'])
prov=[('Milano',1.10,5,5,'Milano Est, Pioltello'),('Torino',1.03,4,4,'Orbassano'),('Genova',1.04,3,3,'Porto, Voltri'),('Bergamo',1.03,4,4,'Orio al Serio'),('Brescia',1.02,4,3,'Brescia Est'),('Novara',1.00,3,3,'CIM Novara'),('Lodi',1.00,3,4,'Lodi, Casalpusterlengo'),('Alessandria',0.97,3,2,'Rivalta Scrivia'),('Pavia',0.97,2,2,'Mortara'),('Aosta',1.00,1,1,'Aosta')]
P0=r+2
for i,(n,c,b,cc,polo) in enumerate(prov,P0):
    put(ws,i,1,n,bold=True); put(ws,i,2,c,'0.00',True); put(ws,i,3,b,'0',True); put(ws,i,4,cc,'0',True); put(ws,i,5,polo); put(ws,i,6,f'=$B$6*B{i}',EUR)
P_END=P0+len(prov)-1

# 3. Welfare
r=P_END+2
section(ws,r,'3. Voci di welfare',8)
header(ws,r+1,['Voce','Costo per dipendente (€/mese)','Punti attrattività','Effetto sul morale','Disponibile dal','Note'])
welf=[('Buoni pasto',110,6,'+5','2002',''),('Mensa aziendale',90,5,'+4','2002','In più costo fisso di 1.500 €/mese per sito'),('Assicurazione sanitaria integrativa',35,7,'+5','2002',''),('Premio di produzione',80,6,'+6','2002','Legato ai KPI del magazzino'),('Orari flessibili',0,5,'+5','2002','Costa 2% di produttività invece di denaro'),('Formazione continua',25,4,'+3','2002','Aumenta anche l\'esperienza'),('Navetta dalla stazione',20,3,'+2','2002','Utile nei poli fuori città'),('Asilo convenzionato',40,4,'+4','2008',''),('Smart working (solo uffici)',0,6,'+6','2020','Solo impiegati e manager')]
W0=r+2
for i,(n,c,pts,m,y,note) in enumerate(welf,W0):
    put(ws,i,1,n,bold=True); put(ws,i,2,c,EUR,True); put(ws,i,3,pts,'0',True); put(ws,i,4,m); put(ws,i,5,y); put(ws,i,6,note,wrapit=True)
W_END=W0+len(welf)-1

# 4. Canali
r=W_END+2
section(ws,r,'4. Canali di ricerca',8)
header(ws,r+1,['Canale','Costo per assunzione (€)','Tempo (giorni)','Qualità media candidati (1-5)','Ruoli','Effetti collaterali'])
ch=[('Annuncio',300,30,2,'Tutti tranne direttori','Qualità molto variabile'),('Agenzia interinale',1200,7,3,'Autisti, magazzinieri, impiegati','Contratto temporaneo, costo orario +15%'),('Cooperativa',0,5,2,'Magazzinieri','Costo orario -10%, morale -10, rischio vertenze'),('Passaparola',0,20,4,'Tutti tranne direttori','Funziona solo con morale aziendale sopra 60'),('Head hunter',15000,45,5,'Manager e direttori','Porta personaggi con tratti rari')]
C0=r+2
for i,row in enumerate(ch,C0):
    put(ws,i,1,row[0],bold=True); put(ws,i,2,row[1],EUR,True); put(ws,i,3,row[2],'0',True); put(ws,i,4,row[3],'0',True); put(ws,i,5,row[4],wrapit=True); put(ws,i,6,row[5],wrapit=True)
C_END=C0+len(ch)-1

# 5. Credibilità
r=C_END+2
section(ws,r,'5. Credibilità aziendale (0-100, partenza 50)',8)
header(ws,r+1,['Evento','Effetto sulla credibilità','Note'])
cred=[('Stipendi pagati puntuali per 12 mesi',5,''),('Apertura di un nuovo magazzino',4,'Segnale di crescita'),('Premio o riconoscimento di settore',6,''),('Welfare migliore della media provinciale',3,'Ogni anno'),('Ritardo nel pagamento degli stipendi',-12,''),('Licenziamenti collettivi',-15,'Dimezzato se ricollochi il personale'),('Incidente grave sul lavoro',-10,''),('Vertenza sindacale',-6,''),('Recensione virale su LavoroBrutto.it',-4,'Evento parodistico')]
K0=r+2
for i,(e,v,n) in enumerate(cred,K0):
    put(ws,i,1,e,bold=True); put(ws,i,2,v,'+0;-0;0',True); put(ws,i,3,n,wrapit=True)
K_END=K0+len(cred)-1

# 6. Calcolatore
r=K_END+2
section(ws,r,'6. Calcolatore: un candidato accetta la nostra offerta?',8)
calc_r=r+1
labels=[('Ruolo',roles[0][0],None,True),('Provincia','Milano',None,True),('Stipendio offerto (RAL €)',25000,EUR,True),('Voci di welfare attive: punti totali',None,'0',False),('Credibilità aziendale (0-100)',50,'0',True),('Prospettive di crescita (0-10)',5,'0',True),('Soglia del candidato (attrattività minima)',55,'0',True)]
rows={}
for k,(lab,val,fmt,isin) in enumerate(labels):
    rr=calc_r+k; rows[lab]=rr
    put(ws,rr,1,lab,bold=True); put(ws,rr,2,val,fmt,isin)
ws.cell(row=rows['Voci di welfare attive: punti totali'],column=2).value=f'=SUM(C{W0}:C{W0+2})'
ws.cell(row=rows['Voci di welfare attive: punti totali'],column=3,value='Esempio: buoni pasto + mensa + assicurazione. Cambia l\'intervallo per provare altre combinazioni.').font=F(color='4A5768',size=9)
rr=calc_r+len(labels)
mk=rr; put(ws,mk,1,'RAL di mercato in provincia (€)',bold=True)
put(ws,mk,2,f'=INDEX($B$6:$B${R_END},MATCH(B{rows["Ruolo"]},$A$6:$A${R_END},0))*INDEX($B${P0}:$B${P_END},MATCH(B{rows["Provincia"]},$A${P0}:$A${P_END},0))',EUR)
cp=mk+1; put(ws,cp,1,'Punti compenso',bold=True); put(ws,cp,2,f'=MAX(-30,MIN(30,(B{rows["Stipendio offerto (RAL €)"]}/B{mk}-1)*200))','0.0')
ws.cell(row=cp,column=3,value='Ogni 1% sopra il mercato = +2 punti (max ±30)').font=F(color='4A5768',size=9)
at=cp+1; put(ws,at,1,'Attrattività totale',bold=True)
put(ws,at,2,f'=30+B{cp}+B{rows["Voci di welfare attive: punti totali"]}+(B{rows["Credibilità aziendale (0-100)"]}-50)*0.5+B{rows["Prospettive di crescita (0-10)"]}','0.0')
ws.cell(row=at,column=3,value='= 30 base + compenso + welfare + (credibilità − 50) × 0,5 + prospettive').font=F(color='4A5768',size=9)
pr=at+1; put(ws,pr,1,'Probabilità che accetti',bold=True)
put(ws,pr,2,f'=MAX(0,MIN(1,0.5+(B{at}-B{rows["Soglia del candidato (attrattività minima)"]})/30))','0%')
ws.cell(row=pr,column=3,value='50% se l\'attrattività è pari alla soglia; ogni 3 punti in più o in meno = ±10%').font=F(color='4A5768',size=9)
ws.cell(row=pr,column=2).font=F(bold=True)
co=pr+1; put(ws,co,1,'Costo azienda annuo dell\'assunzione (€)',bold=True)
put(ws,co,2,f'=B{rows["Stipendio offerto (RAL €)"]}*$H$6+SUM(B{W0}:B{W0+2})*12',EUR)
for c,w in zip('ABCDEFGH',(42,22,24,22,22,46,4,22)): ws.column_dimensions[c].width=w
ws.column_dimensions['G'].width=3

# ================= TECNOLOGIE =================
wt=wb.create_sheet('Tecnologie')
wt['A1']='Investimenti in tecnologia'; wt['A1'].font=title_font
wt['A2']='Ricerca una volta a livello aziendale, poi installazione in ogni magazzino. Effetti in punti percentuali sui KPI del magazzino. Valori di partenza da bilanciare.'; wt['A2'].font=F(color='4A5768')
section(wt,4,'Parametri',15)
params=[('Innovazione del Direttore Tecnologia (0-10)',5,'0'),('Riduzione tempi di ricerca per punto di Innovazione',0.05,'0%'),('Rimborso Bando Fabbrica 4.0 (dal 2017)',0.40,'0%'),('Maggiorazione del leasing sul costo totale',0.18,'0%')]
for k,(lab,v,fmt) in enumerate(params,5):
    put(wt,k,1,lab,bold=True); put(wt,k,2,v,fmt,True)
cols=['Ramo','Tecnologia','Anno di sblocco','Prerequisito','Costo ricerca (€)','Mesi di ricerca (base)','Mesi di ricerca effettivi','Installazione per sito (€)','Canone mensile per sito (€)','Costo primo anno per sito (€)','Effetto puntualità (pp)','Effetto accuratezza (pp)','Settimane di apprendimento','Altri effetti']
header(wt,10,cols)
tech=[
('Informatizzazione','Gestionale di base',2002,'—',15000,2,5000,300,1,1,2,''),
('Informatizzazione','WMS di base',2002,'Gestionale di base',40000,4,25000,800,1,3,6,'Errori di inventario dimezzati'),
('Informatizzazione','EDI con i clienti',2003,'Gestionale di base',30000,3,5000,400,2,1,3,'Richiesto da molti clienti della GDO'),
('Informatizzazione','WMS avanzato multi-sito',2010,'WMS di base',90000,6,40000,1500,1,3,8,'Sposta merce tra magazzini'),
('Magazzino','Lettori barcode',2002,'—',10000,1,12000,150,0,2,2,''),
('Magazzino','Terminali in radiofrequenza',2004,'Lettori barcode; WMS di base',25000,3,30000,400,1,2,4,''),
('Magazzino','Pick-to-light',2008,'Terminali in radiofrequenza',40000,4,80000,600,1,3,4,'Picking +15% più veloce'),
('Magazzino','Magazzino verticale automatico',2012,'WMS avanzato multi-sito',80000,6,400000,2500,1,2,10,'Capacità del sito +20%'),
('Magazzino','Robot mobili (AMR)',2018,'Magazzino verticale automatico',150000,8,600000,4000,1,2,12,'Magazzinieri necessari -30%; tensione sindacale'),
('Trasporto','Telematica GPS',2004,'—',20000,2,10000,500,2,0,3,'Mezzi visibili in tempo reale'),
('Trasporto','Pianificazione delle rotte',2007,'Telematica GPS',50000,5,15000,600,3,0,6,'Gasolio -5%'),
('Trasporto','Mezzi a metano liquido (LNG)',2014,'—',60000,6,50000,300,0,0,2,'Gasolio -20%, credibilità +3'),
('Trasporto','Furgoni elettrici urbani',2020,'Telematica GPS',80000,6,90000,800,0,0,4,'Accesso alle ZTL, credibilità +4'),
('Dati','Cruscotto KPI',2006,'Gestionale di base',20000,2,3000,200,1,1,2,'Avvisi anticipati sui clienti a rischio'),
('Dati','Previsione dei volumi',2013,'Cruscotto KPI; WMS avanzato multi-sito',70000,6,10000,700,2,1,6,'Meno straordinari nei picchi'),
('Dati','Ottimizzazione con IA',2022,'Previsione dei volumi; Pianificazione delle rotte',200000,9,30000,2000,3,2,10,''),
('Persone e sicurezza','Formazione strutturata',2002,'—',8000,1,4000,250,0,1,1,'Morale +3, incidenti -20%'),
('Persone e sicurezza','Sistemi anti-collisione muletti',2010,'Formazione strutturata',15000,2,25000,150,0,1,2,'Incidenti -40%'),
('Persone e sicurezza','Esoscheletri',2019,'Sistemi anti-collisione muletti',40000,4,35000,300,0,1,4,'Assenze per infortunio -25%, morale +3'),
]
T0=11
for i,t in enumerate(tech,T0):
    ramo,nome,anno,pre,cr,mr,inst,can,ep,ea,sett,altri=t
    put(wt,i,1,ramo); put(wt,i,2,nome,bold=True); put(wt,i,3,anno,'0',True); put(wt,i,4,pre,wrapit=True)
    put(wt,i,5,cr,EUR,True); put(wt,i,6,mr,'0',True)
    put(wt,i,7,f'=ROUND(F{i}*MAX(0.4,1-$B$5*$B$6),1)','0.0')
    put(wt,i,8,inst,EUR,True); put(wt,i,9,can,EUR,True)
    put(wt,i,10,f'=H{i}*IF(C{i}>=2017,1-$B$7,1)+I{i}*12',EUR)
    put(wt,i,11,ep,'+0;-0;0',True); put(wt,i,12,ea,'+0;-0;0',True); put(wt,i,13,sett,'0',True); put(wt,i,14,altri,wrapit=True)
T_END=T0+len(tech)-1
n=T_END+2
wt.cell(row=n,column=1,value='Note').font=F(bold=True)
notes=['Durante le settimane di apprendimento il magazzino perde metà dell\'effetto finale invece di guadagnarlo; formazione e morale alto le accorciano.',
'Costo primo anno per sito = installazione (ridotta del rimborso del bando per le tecnologie dal 2017) + 12 canoni.',
'Mesi di ricerca effettivi = mesi base × (1 − Innovazione × riduzione per punto), mai sotto il 40% del tempo base.',
'Con il leasing l\'installazione non si paga subito ma si spalma su 36 mesi, con la maggiorazione indicata nei parametri.']
for k,t in enumerate(notes,n+1):
    c=wt.cell(row=k,column=1,value=t); c.font=F(color='4A5768'); wt.merge_cells(start_row=k,end_row=k,start_column=1,end_column=10)
for c,w in zip('ABCDEFGHIJKLMN',(42,30,10,30,14,11,11,15,14,15,11,11,13,40)): wt.column_dimensions[c].width=w
wt.freeze_panes='C11'
# leasing example
lr=n+len(notes)+2
section(wt,lr,'Esempio leasing',15)
put(wt,lr+1,1,'Tecnologia',bold=True); put(wt,lr+1,2,'Pick-to-light',None,True)
put(wt,lr+2,1,'Installazione per sito (€)',bold=True); put(wt,lr+2,2,f'=INDEX($H${T0}:$H${T_END},MATCH(B{lr+1},$B${T0}:$B${T_END},0))',EUR)
put(wt,lr+3,1,'Rata mensile leasing su 36 mesi (€)',bold=True); put(wt,lr+3,2,f'=B{lr+2}*(1+$B$8)/36',EUR)
put(wt,lr+4,1,'Totale rate + canone per mese (€)',bold=True); put(wt,lr+4,2,f'=B{lr+3}+INDEX($I${T0}:$I${T_END},MATCH(B{lr+1},$B${T0}:$B${T_END},0))',EUR)

wb.move_sheet('Regole',offset=2)
wb.save(P)
print(wb.sheetnames)
